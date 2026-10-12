# -*- coding: utf-8 -*-
"""Backend logic of the Jenkins / DOORS / ALM / QC analysis platform (no GUI dependencies).

Sections
    1. Configuration, paths, persistence
    2. Jenkins API
    3. DOORS CSV reading, scanning and rule analysis (10_AS / 40_Gen_SWXS)
    4. Summaries and cross-reference statistics
    5. ALM (RTC) client
    6. QC (TestInstances / Requirements) analysis
    7. Excel report export
    8. AnalysisBackend facade (full analysis pipeline)

Typical usage
    be = AnalysisBackend()
    be.connect_alm()                              # optional, needs RTC credentials
    status = be.jenkins_status("SomeProject")
    be.load_build("SomeProject", status["last_success_build"])
    req = AnalysisRequest(project=..., build=..., as_file_names=[...], as_features=[...], ...)
    result = be.run_full_analysis(req, progress=print)
    be.export_report("report.xlsx", result, req)
"""

import base64
import csv
import json
import logging
import os
import re
import shutil
import sys
import threading
import time
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Callable, Dict, List, Optional, Tuple

import pandas as pd
import requests
import urllib3
from requests.adapters import HTTPAdapter
from requests.auth import HTTPBasicAuth
from urllib3.util.retry import Retry

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

logger = logging.getLogger(__name__)

try:
    csv.field_size_limit(512 * 1024 * 1024)  # 512MB
except OverflowError:
    pass

LogFunc = Callable[[str, str], None]  # log(message, level)
ProgressFunc = Callable[[int, int, str, int], None]  # progress(step, total_steps, message, percent)


_LOG_LEVELS = {"warning": logging.WARNING, "error": logging.ERROR}


def _noop_log(message, level="info"):
    logger.log(_LOG_LEVELS.get(level, logging.INFO), message)


# ============================================================================
# 1. Configuration, paths, persistence
# ============================================================================
EMPTY_FEATURE = "<EMPTY>"
SEP = "|"
QUOTE = '"'

QC_ENCODING = "utf-8-sig"
QC_CSV_DELIMITER = ";"
FOLDER_COL = "TestSetFolder_Path_full"
RUN_STATUS_COL = "TestExec_Status"

# Jenkins credentials come from JENKINS_USER / JENKINS_TOKEN (backend/.env); never hardcode them.
BASE_URL = "https://rb-jmaas.de.bosch.com/BEG-CN/job/BEG-CN/job/Splunk"
USERNAME = os.environ.get("JENKINS_USER", "")
TOKEN = os.environ.get("JENKINS_TOKEN", "")
AUTH = HTTPBasicAuth(USERNAME, TOKEN)
PROXIES = {
    "http": "http://rb-proxy-apac.bosch.com:8080",
    "https": "http://rb-proxy-apac.bosch.com:8080",
}

# RTC / ALM
SERVER = "https://rb-alm-06-p.de.bosch.com"
PROJECT_AREA = "_UkehAJxnEeeWvYReE95djA"
RTC_HTTP_TIMEOUT = 15
RTC_HTTP_MAX_RETRIES = 2
RTC_HTTP_BACKOFF = 0.3

# Cache layout (same as the GUI so both can share caches)
if getattr(sys, "frozen", False):
    BASE_PATH = os.path.dirname(sys.executable)
else:
    BASE_PATH = os.path.dirname(os.path.abspath(__file__))

APP_CACHE_DIR = os.path.join(BASE_PATH, ".app_cache")
JENKINS_CACHE_DIR = os.path.join(APP_CACHE_DIR, "jenkins")
ALM_CACHE_DIR = os.path.join(APP_CACHE_DIR, "alm")
CONFIG_DIR = os.path.join(APP_CACHE_DIR, "config")
TIMELINE_CACHE_DIR = os.path.join(APP_CACHE_DIR, "timeline")
CRED_DIR = os.path.join(APP_CACHE_DIR, "cred")
for _d in (JENKINS_CACHE_DIR, ALM_CACHE_DIR, CONFIG_DIR, TIMELINE_CACHE_DIR, CRED_DIR):
    os.makedirs(_d, exist_ok=True)

CACHE_DIR = JENKINS_CACHE_DIR
CONFIG_PATH = os.path.join(CONFIG_DIR, "config.json")
ALM_CACHE_PATH = os.path.join(ALM_CACHE_DIR, "alm_cache.json")
TIMELINE_CACHE_PATH = os.path.join(TIMELINE_CACHE_DIR, "timeline_cache.json")
CRED_PATH = os.path.join(CRED_DIR, "rtc_cred.bin")


def load_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, "r") as f:
                return json.load(f)
        except Exception as e:
            logger.warning(f"Configuration load failed, using defaults: {e}")
    return {}


def save_config(config):
    try:
        with open(CONFIG_PATH, "w") as f:
            json.dump(config, f)
    except Exception as e:
        logger.warning(f"Configuration save failed: {e}")


def load_alm_cache():
    if os.path.exists(ALM_CACHE_PATH):
        try:
            with open(ALM_CACHE_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logger.warning(f"ALM cache load failed, will re-fetch: {e}")
    return {}


def save_alm_cache(cache):
    try:
        with open(ALM_CACHE_PATH, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False)
    except Exception as e:
        logger.warning(f"ALM cache save failed: {e}")


def load_local_cache(cache_file=None):
    path = cache_file or TIMELINE_CACHE_PATH
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_local_cache(cache, cache_file=None):
    path = cache_file or TIMELINE_CACHE_PATH
    with open(path, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


# ============================================================================
# 2. Jenkins API
# ============================================================================
_http_session = requests.Session()
try:
    _adapter = HTTPAdapter(max_retries=Retry(
        total=3,
        backoff_factor=0.5,
        status_forcelist=[500, 502, 503, 504],
        allowed_methods=["GET", "POST", "HEAD"],
    ))
    _http_session.mount("https://", _adapter)
    _http_session.mount("http://", _adapter)
except Exception as e:
    logger.warning(f"HTTP retry adapter initialization failed: {e}")

REQ = dict(auth=AUTH, proxies=PROXIES, verify=False, timeout=20)


def get_all_projects():
    r = _http_session.get(f"{BASE_URL}/api/json", **REQ)
    r.raise_for_status()
    return [j["name"] for j in r.json().get("jobs", [])]


def get_last_build_info(project):
    r = _http_session.get(f"{BASE_URL}/job/{project}/lastBuild/api/json", **REQ)
    r.raise_for_status()
    d = r.json()
    if d["building"]:
        return d["number"], "RUNNING"
    return d["number"], d["result"] or "UNKNOWN"


def get_last_success_info(project):
    r = _http_session.get(f"{BASE_URL}/job/{project}/lastSuccessfulBuild/api/json", **REQ)
    r.raise_for_status()
    d = r.json()
    ts = d["timestamp"] / 1000
    return d["number"], datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M")


def get_jenkins_status(project):
    """Last build / last successful build of a project in one call."""
    lb, ls = get_last_build_info(project)
    sb, sd = get_last_success_info(project)
    return {"last_build": lb, "last_status": ls, "last_success_build": sb, "last_success_time": sd}


def trigger_build(project):
    r = _http_session.post(
        f"{BASE_URL}/job/{project}/buildWithParameters",
        data={}, allow_redirects=True, **REQ
    )
    if r.status_code not in (200, 201, 302):
        raise RuntimeError("Trigger build failed")


def trigger_and_wait(project, timeout=40.0, interval=5.0, log: LogFunc = _noop_log):
    """Trigger a build and wait until a new build number appears. Returns jenkins status dict."""
    old_build, _ = get_last_build_info(project)
    trigger_build(project)
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            status = get_jenkins_status(project)
            if status["last_build"] != old_build:
                log(f"Trigger successful, new Build #{status['last_build']}, status: {status['last_status']}", "success")
                return status
        except Exception as e:
            logger.warning(f"Polling Jenkins failed: {e}")
        time.sleep(interval)
    raise TimeoutError("Trigger failed: no new build detected within %d seconds" % int(timeout))


def download_ref_zip(project, build):
    r = _http_session.get(f"{BASE_URL}/job/{project}/{build}/api/json", **REQ)
    r.raise_for_status()
    arts = r.json().get("artifacts", [])
    zips = [a for a in arts if a["fileName"].endswith("_ref.zip")]
    if not zips:
        raise RuntimeError("*_ref.zip not found")
    return zips[0]


def download_file(project, build, artifact):
    r = _http_session.get(
        f"{BASE_URL}/job/{project}/{build}/artifact/{artifact['relativePath']}", **REQ
    )
    r.raise_for_status()
    return r.content


def is_project_build_cached(project, build):
    project = str(project or "").strip()
    build = str(build or "").strip()
    if not project or not build:
        return False
    return os.path.exists(os.path.join(CACHE_DIR, f"{project}_{build}_ref.zip"))


# ============================================================================
# 3. DOORS CSV reading, scanning and rule analysis
# ============================================================================
def read_dumpdoors_csv(path):
    try:
        return pd.read_csv(path, sep=SEP, quotechar=QUOTE, dtype=str, engine="python")
    except Exception as e:
        logger.warning(f"CSV parsing failed: {path}\n{e}")
        return None


def get_external_lookup(module_csv, csv_cache):
    ext_csv = module_csv.replace("-module.csv", "-external.csv")
    lookup_key = ext_csv + "::source_lookup"
    if lookup_key in csv_cache:
        return csv_cache[lookup_key]

    ext_df = csv_cache.get(ext_csv)
    if ext_df is None and ext_csv not in csv_cache:
        if os.path.exists(ext_csv):
            try:
                ext_df = pd.read_csv(ext_csv, sep=SEP, quotechar=QUOTE, dtype=str, engine="python")
            except Exception as e:
                logger.warning(f"Failed to read external CSV {ext_csv}: {e}")
                ext_df = None
        csv_cache[ext_csv] = ext_df

    temp_lookup = {}
    if ext_df is not None and "Source_URL" in ext_df.columns:
        for _, row in ext_df.iterrows():
            source_url = str(row.get("Source_URL", "")).strip()
            if not source_url:
                continue
            if source_url not in temp_lookup:
                temp_lookup[source_url] = {"Names": [], "Description": "", "Target_URL": ""}

            name = str(row.get("Name", "") or "").strip()
            if name and name not in temp_lookup[source_url]["Names"]:
                temp_lookup[source_url]["Names"].append(name)

            # Keep first non-empty value (40_Gen Story info)
            desc = str(row.get("Description", "") or "").strip()
            target = str(row.get("Target_URL", "") or "").strip()
            if not temp_lookup[source_url]["Description"] and desc:
                temp_lookup[source_url]["Description"] = desc
            if not temp_lookup[source_url]["Target_URL"] and target:
                temp_lookup[source_url]["Target_URL"] = target

    lookup = {
        source_url: {
            "Name": "; ".join(info["Names"]),
            "Description": info["Description"],
            "Target_URL": info["Target_URL"],
        }
        for source_url, info in temp_lookup.items()
    }
    csv_cache[lookup_key] = lookup
    return lookup


def get_swxs_source_lookup(csv_cache):
    lookup_key = "__swxs_target_to_source_lookup__"
    if lookup_key in csv_cache:
        return csv_cache[lookup_key]

    lookup = {}
    for source_path, source_df in csv_cache.items():
        if not isinstance(source_path, str) or not source_path.endswith("-links.csv"):
            continue
        if source_df is None or "Target_URL" not in source_df.columns or "Source_URL" not in source_df.columns:
            continue
        for _, row in source_df.iterrows():
            target_url = str(row.get("Target_URL", "")).strip()
            if not target_url:
                continue
            source_url = str(row.get("Source_URL", "") or "").strip()
            if not source_url:
                continue
            if target_url not in lookup:
                lookup[target_url] = []
            if source_url not in lookup[target_url]:
                lookup[target_url].append(source_url)

    # Same Target_URL may map to multiple SWXS: join in scan order
    merged_lookup = {target: "; ".join(sources) for target, sources in lookup.items()}
    csv_cache[lookup_key] = merged_lookup
    return merged_lookup


def _scan_module_files(base_dir, folder_names):
    result = {}
    for root, _, files in os.walk(base_dir):
        if os.path.basename(root) not in folder_names:
            continue
        for f in files:
            if not f.endswith("-module.csv"):
                continue
            module = os.path.join(root, f)
            df = read_dumpdoors_csv(module)
            if df is None or "Modulename" not in df.columns:
                continue
            names = df["Modulename"].dropna().astype(str).str.strip()
            if names.empty:
                continue
            result[names.iloc[0]] = {"module": module}
    return result


def scan_10as_files(base_dir):
    return _scan_module_files(base_dir, ("10_AS",))


def scan_40gen_swxs_files(base_dir):
    return _scan_module_files(base_dir, ("40_Gen_SWXS", "40_Gen_SWSX"))


def scan_40gen_swxs_targets(base_dir, csv_cache):
    targets = set()
    for root, _, files in os.walk(base_dir):
        if os.path.basename(root) not in ("40_Gen_SWXS", "40_Gen_SWSX"):
            continue
        for f in files:
            if not f.endswith("-links.csv"):
                continue
            path = os.path.join(root, f)
            if path not in csv_cache:
                csv_cache[path] = read_dumpdoors_csv(path)
            df = csv_cache[path]
            if df is None or "Target_URL" not in df.columns:
                continue
            targets.update(df["Target_URL"].dropna().astype(str).str.strip())
    return targets


def _filter_module_rows(df, fname, features):
    """Rows of module `fname` whose RB_Baseline is in `features` (EMPTY_FEATURE matches blanks)."""
    df = df[df["Modulename"].astype(str).str.strip() == fname]
    mask = pd.Series(False, index=df.index)
    if EMPTY_FEATURE in features:
        mask |= df["RB_Baseline"].isna() | (df["RB_Baseline"].astype(str).str.strip() == "")
    mask |= df["RB_Baseline"].isin([f for f in features if f != EMPTY_FEATURE])
    return df[mask]


def _is_blank(value):
    return pd.isna(value) or not str(value).strip()


def analyze_file(fname, module_csv, features, csv_cache, xs_targets):
    """10_AS rule check. Returns tuples
    (file, object_text, url, feature, fnid, req_id, reason, alm_name, swxs_source)."""
    if module_csv not in csv_cache:
        csv_cache[module_csv] = read_dumpdoors_csv(module_csv)
    df = csv_cache[module_csv]
    if df is None or df.empty:
        return []
    df = _filter_module_rows(df, fname, features)

    problems = []
    external_lookup = get_external_lookup(module_csv, csv_cache)
    swxs_source_lookup = get_swxs_source_lookup(csv_cache)

    for _, r in df.iterrows():
        status = str(r.get("RB_AS_Status", "")).lower().strip()
        agreed = str(r.get("RB_Agreed_by_Customer", "")).upper().strip()
        change = str(r.get("RB_Change_Required", "")).lower().strip()
        url = str(r.get("URL", "")).strip()
        rb_baseline = r.get("RB_Baseline")
        feature = EMPTY_FEATURE if _is_blank(rb_baseline) else str(rb_baseline)
        obj = str(r.get("Object_Text") or "N/A")
        fnid = str(r.get("RB_Top_FNID", "")).strip()

        def add(reason):
            problems.append((
                fname, obj, url, feature, fnid,
                r.get("Object_Identifier", ""),
                reason,
                external_lookup.get(url, {}).get("Name", ""),
                swxs_source_lookup.get(url, ""),
            ))

        # Independent rules: one requirement can trigger multiple problems
        valid_status = ("accepted", "rejected", "no_req", "validated", "accepted_w_rem", "canceled")
        if status not in valid_status:
            add("Requirement not analyzed or analysis complete")
        if status == "rejected" and agreed != "TRUE":
            add("Rejected requirement not agreed by customer")
        if status in ("accepted", "accepted_w_rem", "validated") and change not in ("yes", "no"):
            add("Change required is empty")
        if status not in ("no_req", "canceled", "rejected") and not fnid:
            add("FNID not assigned")
        if status not in ("no_req", "canceled", "rejected") and _is_blank(rb_baseline):
            add("Baseline not assigned")

        if status in ("accepted", "accepted_w_rem", "validated") and change == "yes":
            if not url:
                add("URL is empty")
            else:
                if url not in xs_targets:
                    add("Not linked to XS")
                if url not in external_lookup:
                    add("ALM not assigned")

    return problems


def analyze_40gen_swxs_file(fname, module_csv, features, csv_cache):
    """40_Gen_SWXS rule check. Returns one tuple per requirement
    (file, object_text, url, feature, reasons, story_desc, story_url, rs_status, test_status, review_status)."""
    if module_csv not in csv_cache:
        csv_cache[module_csv] = read_dumpdoors_csv(module_csv)
    df = csv_cache[module_csv]
    if df is None or df.empty:
        return []
    df = _filter_module_rows(df, fname, features)

    external_lookup = get_external_lookup(module_csv, csv_cache)
    problems = []

    for _, r in df.iterrows():
        status = str(r.get("RB_RS_CP_Status", "")).lower().strip()
        test_status = str(r.get("RB_TestStatus", "")).lower().strip()
        review_status = str(r.get("RB_ReviewStatus", "")).lower().strip()
        url = str(r.get("URL", "")).strip()
        rb_baseline = r.get("RB_Baseline")
        obj = str(r.get("Object_Text") or "N/A")
        feature = EMPTY_FEATURE if _is_blank(rb_baseline) else str(rb_baseline)

        story_desc = ""
        story_url = ""
        has_story = False
        if url:
            ext_info = external_lookup.get(url)
            if ext_info:
                has_story = True
                story_desc = ext_info.get("Description", "")
                story_url = ext_info.get("Target_URL", "")

        reasons = []
        if status in ("accepted", "validated"):
            if test_status != "passed":
                reasons.append("XS not test passed")
            if review_status != "ok":
                reasons.append("XS requirement review not completed")
            if not has_story:
                reasons.append("XS requirement not linked to Story")

        problems.append((fname, obj, url, feature, "; ".join(reasons), story_desc, story_url,
                         status, test_status, review_status))

    return problems


def sort_problem_first(rows):
    """Rows with a non-empty 'Issue Reason' first (stable)."""
    return sorted(rows, key=lambda r: 0 if str(r.get("Issue Reason", "")).strip() else 1)


class DoorsDataSession:
    """State of one downloaded Jenkins build: scanned files, CSV cache, link targets."""

    def __init__(self):
        self.file_map = {}       # 10_AS: module name -> {"module": csv path}
        self.xs_file_map = {}    # 40_Gen_SWXS
        self.csv_cache = {}
        self.xs_targets = set()
        self.base_dir = None
        self.project = None
        self.build = None

    def clear(self):
        self.file_map = {}
        self.xs_file_map = {}
        self.csv_cache.clear()
        self.xs_targets = set()
        self.base_dir = None
        self.project = None
        self.build = None

    @property
    def loaded(self):
        return bool(self.file_map and self.xs_file_map and self.base_dir)

    def download_and_scan(self, project, build, log: LogFunc = _noop_log):
        """Download (or reuse cached) *_ref.zip, extract, scan and drop older builds of the project.
        Returns (sorted 10_AS names, sorted 40_Gen_SWXS names)."""
        build = str(build).strip()

        cfg = load_config()
        cfg["last_project"] = project
        save_config(cfg)

        zip_cache_path = os.path.join(CACHE_DIR, f"{project}_{build}_ref.zip")
        self.base_dir = os.path.join(CACHE_DIR, f"{project}_{build}_ref")
        os.makedirs(self.base_dir, exist_ok=True)

        if os.path.exists(zip_cache_path):
            log("Using local cached ZIP", "info")
        else:
            log("Downloading Jenkins artifact...", "info")
            artifact = download_ref_zip(project, build)
            data = download_file(project, build, artifact)
            with open(zip_cache_path, "wb") as f:
                f.write(data)

        log("Extracting ZIP...", "info")
        with zipfile.ZipFile(zip_cache_path) as z:
            z.extractall(self.base_dir)

        log("Scanning 10_AS files...", "info")
        self.file_map = scan_10as_files(self.base_dir)
        log("Scanning 40_Gen_SWXS files...", "info")
        self.xs_file_map = scan_40gen_swxs_files(self.base_dir)
        self.csv_cache.clear()
        self.xs_targets = scan_40gen_swxs_targets(self.base_dir, self.csv_cache)

        # Keep only the current build of this project in the cache dir
        current_zip = f"{project}_{build}_ref.zip"
        current_dir = f"{project}_{build}_ref"
        removed = 0
        for name in os.listdir(CACHE_DIR):
            if name in (current_zip, current_dir):
                continue
            if name.startswith(f"{project}_") and (name.endswith("_ref.zip") or name.endswith("_ref")):
                old_path = os.path.join(CACHE_DIR, name)
                try:
                    if os.path.isdir(old_path):
                        shutil.rmtree(old_path, ignore_errors=True)
                    else:
                        os.remove(old_path)
                    removed += 1
                except Exception as e:
                    log(f"Failed to clean old cache {name}: {e}", "warning")
        if removed:
            log(f"Cleaned {removed} old build cache files/directories", "info")

        cfg["last_build"] = build
        save_config(cfg)
        self.project = project
        self.build = build
        return sorted(self.file_map.keys()), sorted(self.xs_file_map.keys())

    def _module_df(self, module):
        if module not in self.csv_cache:
            self.csv_cache[module] = read_dumpdoors_csv(module)
        return self.csv_cache[module]

    def list_baselines(self, file_names, kind="as"):
        """RB_Baseline values available in the selected files (EMPTY_FEATURE first if blanks exist)."""
        fmap = self.file_map if kind == "as" else self.xs_file_map
        baselines = set()
        has_empty = False
        for fname in file_names:
            df = self._module_df(fmap[fname]["module"])
            if df is None or "RB_Baseline" not in df.columns:
                continue
            col = df["RB_Baseline"]
            if col.isna().any() or (col.astype(str).str.strip() == "").any():
                has_empty = True
            for v in col.dropna().astype(str):
                v = v.strip()
                if v:
                    baselines.add(v)
        return ([EMPTY_FEATURE] if has_empty else []) + sorted(baselines)

    def analyze_as(self, file_names, features):
        """10_AS analysis. Returns all rows (dict) in file order; problems carry 'Issue Reason'."""
        all_rows = []
        problem_map = {}

        for fname in file_names:
            module = self.file_map[fname]["module"]
            df = self._module_df(module)
            if df is None or df.empty:
                continue
            df = _filter_module_rows(df, fname, features)
            external_lookup = get_external_lookup(module, self.csv_cache)
            swxs_source_lookup = get_swxs_source_lookup(self.csv_cache)

            for _, row in df.iterrows():
                rb_baseline = row.get("RB_Baseline")
                url = str(row.get("URL", "")).strip()
                all_rows.append((
                    fname,
                    str(row.get("Object_Text") or "N/A"),
                    url,
                    EMPTY_FEATURE if _is_blank(rb_baseline) else str(rb_baseline),
                    str(row.get("RB_Top_FNID", "")).strip(),
                    str(row.get("Object_Identifier", "")).strip(),
                    str(row.get("RB_AS_Status", "")).strip(),
                    str(row.get("RB_Change_Required", "")).strip(),
                    str(row.get("RB_Agreed_by_Customer", "")).strip(),
                    external_lookup.get(url, {}).get("Name", ""),
                    swxs_source_lookup.get(url, ""),
                    str(row.get("RB_Planning_Feature", "")).strip(),
                    str(row.get("RB_Allocation", "")).strip(),
                ))

            for res in analyze_file(fname, module, features, self.csv_cache, self.xs_targets):
                fname2, obj2, url2, feat2, fnid2, reqid2, reason, alm2, swxs2 = res
                key = (fname2, obj2, url2, feat2, fnid2, reqid2, alm2, swxs2)
                problem_map.setdefault(key, set()).add(reason)

        rows = []
        for (fname, obj, url, feature, fnid, reqid, rb_s, rb_c, rb_a, alm_name, swxs, rb_pf, rb_alloc) in all_rows:
            reasons = problem_map.get((fname, obj, url, feature, fnid, reqid, alm_name, swxs), set())
            merged = "; ".join(sorted(x for x in reasons if x and str(x).strip()))
            rows.append({
                "File_name": fname, "Object_Text": obj, "URL": url,
                "RB_Baseline": feature, "RB_Top_FNID": fnid, "Object_Identifier": reqid,
                "RB_AS_Status": rb_s, "RB_Change_Required": rb_c, "RB_Agreed_by_Customer": rb_a,
                "Issue Reason": merged,
                "Corresponding ALM": alm_name,
                "Corresponding SWXS": swxs,
                "RB_Planning_Feature": rb_pf,
                "RB_Allocation": rb_alloc,
            })
        return rows

    def analyze_xs(self, file_names, features):
        """40_Gen_SWXS analysis. Returns rows (dict), one per requirement."""
        results = []
        for fname in file_names:
            module = self.xs_file_map[fname]["module"]
            results.extend(analyze_40gen_swxs_file(fname, module, features, self.csv_cache))

        merged_map = {}
        for (fname, obj, url, feature, reason, story_desc, story_url, rb_rs, rb_ts, rb_rv) in results:
            key = (fname, obj, url, feature)
            if key not in merged_map:
                merged_map[key] = {
                    "reasons": set(), "story_desc": story_desc, "story_url": story_url,
                    "RB_RS_CP_Status": rb_rs, "RB_TestStatus": rb_ts, "RB_ReviewStatus": rb_rv,
                }
            if reason and str(reason).strip():
                merged_map[key]["reasons"].add(str(reason).strip())

        rows = []
        for (fname, obj, url, feature), info in merged_map.items():
            rows.append({
                "File_name": fname, "Object_Text": obj, "URL": url, "RB_Baseline": feature,
                "RB_RS_CP_Status": info["RB_RS_CP_Status"],
                "RB_TestStatus": info["RB_TestStatus"],
                "RB_ReviewStatus": info["RB_ReviewStatus"],
                "Issue Reason": "; ".join(sorted(info["reasons"])),
                "Story Description": info["story_desc"],
                "Story_URL": info["story_url"],
            })
        return rows


# ============================================================================
# 4. Summaries and cross-reference statistics
# ============================================================================
def _build_issue_summary(title, rows, selected_files=None, features=None):
    total = len(rows)
    problem_count = sum(1 for r in rows if str(r.get("Issue Reason", "")).strip())

    file_stats = {}
    for r in rows:
        st = file_stats.setdefault(r.get("File_name", ""), {"total": 0, "problem": 0, "reasons": {}})
        st["total"] += 1
        reasons_raw = str(r.get("Issue Reason", "")).strip()
        if reasons_raw:
            st["problem"] += 1
            for reason in [x.strip() for x in reasons_raw.split(";") if x.strip()]:
                st["reasons"][reason] = st["reasons"].get(reason, 0) + 1

    global_reasons = {}
    for st in file_stats.values():
        for reason, count in st["reasons"].items():
            global_reasons[reason] = global_reasons.get(reason, 0) + count

    lines = [title]
    parts = [f"Total: {total}", f"Issues: {problem_count}"]
    if selected_files is not None:
        parts.append(f"Files: {len(selected_files)}")
    if features is not None:
        parts.append(f"Baselines: {len(features)}")
    parts.append(f"Files involved: {len(file_stats)}")
    lines.append("  |  ".join(parts))

    if global_reasons:
        reason_strs = [f"{r}: {c}" for r, c in sorted(global_reasons.items(), key=lambda x: (-x[1], x[0]))]
        lines.append("Issues: " + "  |  ".join(reason_strs))
    else:
        lines.append("Issues: None")
    lines.append("")

    lines.append("By file:")
    for fname in sorted(file_stats):
        st = file_stats[fname]
        seg = f"{fname}: total {st['total']}/issues {st['problem']}"
        if st["reasons"]:
            reason_strs = [f"{r}:{c}" for r, c in sorted(st["reasons"].items(), key=lambda x: (-x[1], x[0]))]
            seg += "  [" + ", ".join(reason_strs) + "]"
        else:
            seg += "  [no issues]"
        lines.append("  " + seg)
    return lines


def build_as_summary(rows, selected_files=None, features=None):
    return _build_issue_summary("10_AS Analysis Summary", rows, selected_files, features)


def build_xs_summary(rows, selected_files=None, features=None):
    return _build_issue_summary("40_Gen_SWXS Analysis Summary", rows, selected_files, features)


# Old caches used Chinese reason texts
_REASON_CN_MAP = {
    "Task未完全关闭": "Task not fully closed",
    "AWM评估未完成": "AWM assessment not completed",
    "未完成AWM评估": "AWM assessment not completed",
}

_AWM_APPROVAL_STRS = (
    "Compliance with Product Compliance Code (MP-105.01) approved",
    "reference to Codex Concern (formely known as AWM) provided | Approved",
)


def build_alm_summary_text(all_defects, qc_issues):
    total_count = len(all_defects)
    issue_count = len(qc_issues)
    normal_count = total_count - issue_count

    type_counter = {}
    for d in all_defects:
        t = str(d.get("Type", "Unknown") or "Unknown")
        type_counter[t] = type_counter.get(t, 0) + 1

    reason_counter = {}
    for d in qc_issues:
        reasons_raw = str(d.get("Issue Reason") or d.get("问题原因") or "Unknown Reason")
        reasons = [r.strip() for r in reasons_raw.split(";") if r.strip()] or ["Unknown Reason"]
        for r in reasons:
            r = _REASON_CN_MAP.get(r, r)
            reason_counter[r] = reason_counter.get(r, 0) + 1

    lines = ["=" * 80, " ALM Data Statistics Summary", "=" * 80,
             f"Total data: {total_count}", f"Normal data: {normal_count}", f"Issue data: {issue_count}",
             "-" * 80]
    if type_counter:
        lines.append(" Data type distribution:")
        for type_name, count in sorted(type_counter.items(), key=lambda x: x[1], reverse=True):
            pct = (count / total_count * 100) if total_count else 0
            lines.append(f"  {type_name:<20}: {count:>4} ({pct:>6.1f}%)")
    if reason_counter:
        lines.append("-" * 80)
        lines.append(" Issue Reason Distribution: ")
        for reason, count in sorted(reason_counter.items(), key=lambda x: x[1], reverse=True):
            pct = (count / issue_count * 100) if issue_count else 0
            lines.append(f"  {reason:<30}: {count:>4} ({pct:>6.1f}%)")
    lines.append("=" * 80)
    return "\n".join(lines) + "\n"


def _pct(n, d):
    return f"{n / d * 100:.1f}%" if d else "—"


def build_cross_ref_stats(all_as_rows, all_xs_results, alm_defects):
    """Cross-reference statistics. Returns dict with keys
    as_status, alm_xref, xs_xref, xs_status, alm_status."""
    # 1. 10_AS status distribution
    groups_def = {
        "Accepted\n(accepted/accepted_w_rem/validated/closed)": {"accepted", "accepted_w_rem", "validated", "closed"},
        "Rejected": {"rejected"},
        "No requirement/Canceled\n(no_req/canceled)": {"no_req", "canceled"},
        "Pending analysis\n(new/in_analysis/ext_analysis/modified)": {"new", "in_analysis", "ext_analysis", "modified"},
    }
    as_group_counts = {k: 0 for k in groups_def}
    as_other = 0
    for row in all_as_rows:
        st = str(row.get("RB_AS_Status", "")).lower().strip()
        for grp, valid in groups_def.items():
            if st in valid:
                as_group_counts[grp] += 1
                break
        else:
            as_other += 1

    accepted_set = {"accepted", "accepted_w_rem", "validated", "closed"}
    accepted_rows = [r for r in all_as_rows if str(r.get("RB_AS_Status", "")).lower().strip() in accepted_set]
    accepted_change_yes = [r for r in accepted_rows if str(r.get("RB_Change_Required", "")).lower().strip() == "yes"]
    acc_yes_total = len(accepted_change_yes)
    acc_no_total = len(accepted_rows) - acc_yes_total
    acc_yes_alm_has = sum(1 for r in accepted_change_yes if str(r.get("Corresponding ALM", "")).strip())
    acc_yes_swxs_has = sum(1 for r in accepted_change_yes if str(r.get("Corresponding SWXS", "")).strip())

    # 2. 10_AS <-> ALM Epic
    need_alm_rows = [
        r for r in all_as_rows
        if str(r.get("RB_AS_Status", "")).lower().strip() in ("accepted", "accepted_w_rem", "validated")
        and str(r.get("RB_Change_Required", "")).lower().strip() == "yes"
    ]
    alm_epic_ids = set()
    for d in (alm_defects or []):
        if str(d.get("Type", "")).strip().lower() == "epic":
            aid = str(d.get("ID", "")).strip()
            if aid:
                alm_epic_ids.add(aid)

    def extract_alm_ids(alm_str):
        """'6997105: Title; 6997106: Title2' -> ['6997105', '6997106']"""
        ids = []
        for part in alm_str.split(";"):
            part = part.strip()
            colon_pos = part.find(":")
            if colon_pos > 0 and part[:colon_pos].strip():
                ids.append(part[:colon_pos].strip())
        return ids

    referenced_alm_ids = set()
    need_alm_has_str = need_alm_matched = need_alm_unmatched = 0
    for row in need_alm_rows:
        alm_str = str(row.get("Corresponding ALM", "")).strip()
        if not alm_str:
            continue
        need_alm_has_str += 1
        ids = extract_alm_ids(alm_str)
        referenced_alm_ids.update(ids)
        if any(i in alm_epic_ids for i in ids):
            need_alm_matched += 1
        else:
            need_alm_unmatched += 1

    # 3. 10_AS <-> 40_Gen_SWXS (accepted/validated only)
    xs_accepted_rows = [
        r for r in all_xs_results
        if str(r.get("RB_RS_CP_Status", "")).lower().strip() in ("accepted", "validated")
    ]
    xs_accepted_urls = {str(r.get("URL", "")).strip() for r in xs_accepted_rows if str(r.get("URL", "")).strip()}
    as_with_swxs_rows = [r for r in all_as_rows if str(r.get("Corresponding SWXS", "")).strip()]

    referenced_xs_urls = set()
    xs_matched = xs_unmatched = 0
    for row in as_with_swxs_rows:
        urls_in_row = [u.strip() for u in str(row.get("Corresponding SWXS", "")).split(";") if u.strip()]
        referenced_xs_urls.update(urls_in_row)
        if any(u in xs_accepted_urls for u in urls_in_row):
            xs_matched += 1
        else:
            xs_unmatched += 1

    # 4. 40_Gen_SWXS status distribution
    xs_total = len(all_xs_results)
    xs_accepted_cnt = len(xs_accepted_rows)
    xs_acc_story_has = sum(1 for r in xs_accepted_rows if str(r.get("Story_URL", "")).strip())
    xs_acc_test_pass = sum(1 for r in xs_accepted_rows if str(r.get("RB_TestStatus", "")).lower().strip() == "passed")
    xs_acc_review_ok = sum(1 for r in xs_accepted_rows if str(r.get("RB_ReviewStatus", "")).lower().strip() == "ok")

    # 5. ALM statistics by Type
    alm_total = len(alm_defects or [])
    alm_closed = 0
    alm_type_counts = {}
    alm_type_order = []
    alm_type_awm = {}
    for d in (alm_defects or []):
        if str(d.get("State", "")).strip().lower() == "closed":
            alm_closed += 1
        t = str(d.get("Type", "")).strip() or "Unknown"
        if t not in alm_type_counts:
            alm_type_counts[t] = 0
            alm_type_order.append(t)
        alm_type_counts[t] += 1
        if t in ("Epic", "Defect Fix"):
            awm = alm_type_awm.setdefault(t, {"done": 0, "not_done": 0})
            approvals = str(d.get("Approvals", ""))
            awm["done" if any(s in approvals for s in _AWM_APPROVAL_STRS) else "not_done"] += 1
    alm_not_closed = alm_total - alm_closed

    return {
        "as_status": {
            "total": len(all_as_rows),
            "groups": as_group_counts,
            "other": as_other,
            "accepted_total": len(accepted_rows),
            "accepted_change_yes": acc_yes_total,
            "accepted_change_no": acc_no_total,
            "acc_yes_alm_has": acc_yes_alm_has,
            "acc_yes_alm_none": acc_yes_total - acc_yes_alm_has,
            "acc_yes_alm_has_pct": _pct(acc_yes_alm_has, acc_yes_total),
            "acc_yes_alm_none_pct": _pct(acc_yes_total - acc_yes_alm_has, acc_yes_total),
            "acc_yes_swxs_has": acc_yes_swxs_has,
            "acc_yes_swxs_none": acc_yes_total - acc_yes_swxs_has,
            "acc_yes_swxs_has_pct": _pct(acc_yes_swxs_has, acc_yes_total),
            "acc_yes_swxs_none_pct": _pct(acc_yes_total - acc_yes_swxs_has, acc_yes_total),
        },
        "alm_xref": {
            "need_alm_total": len(need_alm_rows),
            "need_alm_has_str": need_alm_has_str,
            "need_alm_no_str": len(need_alm_rows) - need_alm_has_str,
            "need_alm_matched": need_alm_matched,
            "need_alm_unmatched": need_alm_unmatched,
            "alm_epic_total": len(alm_epic_ids),
            "alm_epic_with_as": len(alm_epic_ids & referenced_alm_ids),
            "alm_epic_without_as": len(alm_epic_ids - referenced_alm_ids),
        },
        "xs_xref": {
            "as_with_swxs": len(as_with_swxs_rows),
            "xs_matched": xs_matched,
            "xs_unmatched": xs_unmatched,
            "xs_accepted_total": len(xs_accepted_urls),
            "xs_with_as": len(xs_accepted_urls & referenced_xs_urls),
            "xs_without_as": len(xs_accepted_urls - referenced_xs_urls),
        },
        "xs_status": {
            "total": xs_total,
            "accepted_validated": xs_accepted_cnt,
            "other": xs_total - xs_accepted_cnt,
            "story_has": xs_acc_story_has,
            "story_none": xs_accepted_cnt - xs_acc_story_has,
            "story_has_pct": _pct(xs_acc_story_has, xs_accepted_cnt),
            "story_none_pct": _pct(xs_accepted_cnt - xs_acc_story_has, xs_accepted_cnt),
            "test_pass": xs_acc_test_pass,
            "test_fail": xs_accepted_cnt - xs_acc_test_pass,
            "test_pass_pct": _pct(xs_acc_test_pass, xs_accepted_cnt),
            "test_fail_pct": _pct(xs_accepted_cnt - xs_acc_test_pass, xs_accepted_cnt),
            "review_ok": xs_acc_review_ok,
            "review_no": xs_accepted_cnt - xs_acc_review_ok,
            "review_ok_pct": _pct(xs_acc_review_ok, xs_accepted_cnt),
            "review_no_pct": _pct(xs_accepted_cnt - xs_acc_review_ok, xs_accepted_cnt),
        },
        "alm_status": {
            "total": alm_total,
            "closed": alm_closed,
            "closed_pct": _pct(alm_closed, alm_total),
            "not_closed": alm_not_closed,
            "not_closed_pct": _pct(alm_not_closed, alm_total),
            "type_counts": alm_type_counts,
            "type_order": alm_type_order,
            "type_awm": alm_type_awm,
        },
    }


# ============================================================================
# 5. ALM (RTC) client
# ============================================================================
OSCL_QUERY = "oslc_cm.query"
OSCL_PROPS = "oslc_cm.properties"
PROP_ID = "dc:identifier"
PROP_TITLE = "dc:title"
PROP_TYPE = "dc:type"
PROP_STATE = "rtc_cm:state"
PROP_RESOLUTION = "rtc_cm:resolution"
PROP_SEVERITY = "oslc_cm:severity"
PROP_FILED_AGAINST = "rtc_cm:filedAgainst"
HEADERS_JSON = {"Accept": "application/json"}

_oslc_cache = {}
_oslc_cache_lock = threading.Lock()
_approval_cache = {}
_approval_cache_lock = threading.Lock()


class RTCAuthError(RuntimeError):
    """RTC credentials missing or rejected."""


# ---- credentials (stored base64 encoded, same file format as the GUI) ----
def encrypt_cred(username, password):
    return base64.b64encode(f"{username}\n{password}".encode("utf-8"))


def decrypt_cred(b64):
    try:
        lines = base64.b64decode(b64).decode("utf-8").split("\n")
        if len(lines) == 2:
            return lines[0], lines[1]
    except Exception:
        pass
    return None, None


def save_cred(username, password, cred_file=None):
    with open(cred_file or CRED_PATH, "wb") as f:
        f.write(encrypt_cred(username, password))


def load_cred(cred_file=None):
    path = cred_file or CRED_PATH
    if not os.path.exists(path):
        return None, None
    with open(path, "rb") as f:
        return decrypt_cred(f.read())


def _make_insecure_session():
    """Session with SSL verification disabled and system proxy settings ignored."""
    s = requests.Session()
    s.verify = False
    s.trust_env = False

    class InsecureHTTPSAdapter(HTTPAdapter):
        def init_poolmanager(self, *args, **kwargs):
            import ssl
            ctx = ssl.create_default_context()
            ctx.check_hostname = False
            ctx.verify_mode = ssl.CERT_NONE
            kwargs["ssl_context"] = ctx
            super().init_poolmanager(*args, **kwargs)

    s.mount("https://", InsecureHTTPSAdapter())
    return s


def check_rtc_login(username, password):
    """Returns (True, "") on success, (False, "AUTH_FAILED") on wrong credentials,
    (False, message) on network errors."""
    try:
        s = _make_insecure_session()
        s.auth = (username, password)
        s.get(f"{SERVER}/ccm/authenticated/identity", timeout=30, verify=False)
        r = s.post(f"{SERVER}/ccm/j_security_check",
                   data={"j_username": username, "j_password": password},
                   timeout=30, verify=False)
        if "x-com-ibm-team-repository-web-auth-msg" in r.headers:
            return False, "AUTH_FAILED"
        return True, ""
    except requests.exceptions.ConnectTimeout:
        return False, "Connection timeout — server unreachable, please check VPN/network"
    except requests.exceptions.ReadTimeout:
        return False, "Read timeout — server response too slow"
    except requests.exceptions.ProxyError as e:
        return False, f"Proxy error — please check proxy settings:\n{e}"
    except requests.exceptions.SSLError as e:
        return False, f"SSL error:\n{e}"
    except requests.exceptions.ConnectionError as e:
        return False, f"Connection failed — possibly DNS resolution failure or network unreachable:\n{e}"
    except Exception as e:
        return False, f"{type(e).__name__}: {e}"


def resolve_rtc_credentials(username=None, password=None, cred_file=None):
    """Explicit credentials (saved on success) or the stored ones. Raises RTCAuthError."""
    explicit = bool(username and password)
    if not explicit:
        username, password = load_cred(cred_file)
    if not username or not password:
        raise RTCAuthError("No RTC credentials available: pass username/password once to store them")
    ok, err = check_rtc_login(username, password)
    if not ok:
        if err.startswith("AUTH_FAILED"):
            raise RTCAuthError("Username or password is incorrect")
        raise RTCAuthError(f"Unable to connect to RTC server: {err}")
    if explicit:
        save_cred(username, password, cred_file or CRED_PATH)
    return username, password


def rtc_setup_session(username=None, password=None, cred_file=None):
    """Authenticated requests.Session for RTC. Raises RTCAuthError."""
    logger.info("Initializing RTC session...")
    username, password = resolve_rtc_credentials(username, password, cred_file or CRED_PATH)
    s = _make_insecure_session()
    s.auth = (username, password)

    try:
        adapter = HTTPAdapter(max_retries=Retry(
            total=RTC_HTTP_MAX_RETRIES,
            backoff_factor=RTC_HTTP_BACKOFF,
            status_forcelist=[500, 502, 503, 504],
            allowed_methods=["GET", "POST", "HEAD"],
        ))
        s.mount("https://", adapter)
        s.mount("http://", adapter)
    except Exception as e:
        logger.warning(f"Retry adapter initialization failed, will use no-retry mode: {e}")

    s.get(f"{SERVER}/ccm/authenticated/identity", timeout=RTC_HTTP_TIMEOUT, verify=False)
    r = s.post(f"{SERVER}/ccm/j_security_check",
               data={"j_username": username, "j_password": password},
               timeout=RTC_HTTP_TIMEOUT, verify=False)
    if "x-com-ibm-team-repository-web-auth-msg" in r.headers:
        raise RTCAuthError("RTC authentication failed")
    logger.info("RTC session established successfully")
    return s


# ---- OSLC helpers ----
def xml_root(text):
    return ET.fromstring(text)


def xml_text(elem, tag):
    return elem.findtext(f"{{*}}{tag}", default="")


def get_oslc_results(j):
    for k in ("oslc:results", "oslc_cm:results", "rdfs:member"):
        if k in j:
            return j[k]
    return []


def resolve_simple_attr(session, uri):
    if not uri:
        return ""
    with _oslc_cache_lock:
        if uri in _oslc_cache:
            return _oslc_cache[uri]
    try:
        if "/users/" in uri:
            val = uri.rsplit("/", 1)[-1]
        else:
            r = session.get(uri, headers=HEADERS_JSON, timeout=RTC_HTTP_TIMEOUT, verify=False)
            val = r.json().get("dc:title", "")
    except Exception as e:
        logger.warning(f"Failed to resolve attribute {uri}: {e}")
        val = ""
    with _oslc_cache_lock:
        _oslc_cache[uri] = val
    return val


def resolve_approvals_via_rpt(session, workitem_id):
    if not workitem_id:
        return ""
    with _approval_cache_lock:
        if workitem_id in _approval_cache:
            return _approval_cache[workitem_id]

    url = f"{SERVER}/ccm/rpt/repository/workitem"
    params = {
        "fields": f"workitem/workItem[id={workitem_id}]/approvalDescriptors/(name|cumulativeStateName)",
        "size": 100,
    }
    approvals = []
    try:
        r = session.get(url, params=params, timeout=RTC_HTTP_TIMEOUT, verify=False)
        root = ET.fromstring(r.text)
        for ad in root.findall(".//approvalDescriptors"):
            name = ad.findtext("name", "").strip()
            state = ad.findtext("cumulativeStateName", "").strip()
            if name or state:
                approvals.append(f"{name} | {state}")
    except Exception as e:
        logger.warning(f"Failed to get approval info {workitem_id}: {e}")

    result = "; ".join(approvals)
    with _approval_cache_lock:
        _approval_cache[workitem_id] = result
    return result


def rtc_query_defects_huge(session, planned_for_id, progress_cb: Optional[Callable[[str], None]] = None):
    """All work items planned for an iteration (keyset pagination on dc:identifier)."""
    logger.info(f"Start querying defects, planned_for_id: {planned_for_id}")
    url = f"{SERVER}/ccm/oslc/contexts/{PROJECT_AREA}/workitems"
    props = ",".join([PROP_ID, PROP_TITLE, PROP_TYPE, PROP_STATE, PROP_RESOLUTION,
                      PROP_SEVERITY, PROP_FILED_AGAINST, "rtc_cm:ownedBy"])

    merged = []
    last_id = None
    page_count = 0
    while True:
        query = f'rtc_cm:plannedFor="{planned_for_id}"'
        if last_id:
            query += f' and dc:identifier>"{last_id}"'
        query += " /sort=dc:identifier"
        params = {OSCL_QUERY: query, OSCL_PROPS: props, "oslc.pageSize": 500}
        try:
            r = session.get(url, params=params, headers=HEADERS_JSON, timeout=30, verify=False)
            r.raise_for_status()
            results = get_oslc_results(r.json())
            if not results:
                break
            merged.extend(results)
            page_count += 1
            last_id = results[-1].get(PROP_ID)
            if progress_cb:
                progress_cb(f"Fetched page {page_count}, total {len(merged)} records")
        except Exception as e:
            logger.warning(f"Failed to query page {page_count + 1}: {e}")
            break

    logger.info(f"Query complete, total {len(merged)} records")
    return merged


def alm_quality_check(defects):
    """Rule check on parsed ALM work items. Returns copies of failing items with 'Issue Reason'."""
    issues = []
    for d in defects:
        reasons = []
        res = d["Resolution"]
        if res not in ("Duplicate", "Invalid"):
            if d["State"] != "Closed":
                reasons.append("Task not fully closed")
            if d["Type"] in ("Epic", "Defect Fix"):
                approvals = d["Approvals"]
                if not any(s in approvals for s in _AWM_APPROVAL_STRS):
                    reasons.append("AWM assessment not completed")
        if reasons:
            d_copy = d.copy()
            d_copy["Issue Reason"] = "; ".join(reasons)
            issues.append(d_copy)
    return issues


def flatten_alm_scope_options(cache=None):
    """Leaf iterations of the timeline cache as
    [{"display", "timeline", "path", "label", "url"}, ...] sorted by display."""
    cache = load_local_cache(TIMELINE_CACHE_PATH) if cache is None else cache
    options = []

    def walk(timeline, nodes, chain, visited_refs=None):
        visited_refs = visited_refs or set()
        for node in nodes or []:
            label = str(node.get("label") or "").strip()
            url = str(node.get("url") or "").strip()
            children = node.get("children") or []
            current_chain = chain + [label]
            ref_nodes = []
            if not children and label and label in cache and label != timeline and label not in visited_refs:
                ref_nodes = cache.get(label) or []
            if ref_nodes:
                walk(timeline, ref_nodes, current_chain, visited_refs | {label})
            elif url and not children:
                path = " / ".join(current_chain)
                options.append({
                    "display": f"{timeline} | {path}",
                    "timeline": timeline,
                    "path": path,
                    "label": label,
                    "url": url,
                })
            else:
                walk(timeline, children, current_chain, visited_refs)

    for timeline, nodes in cache.items():
        walk(timeline, nodes, [])
    return sorted(options, key=lambda item: item["display"].lower())


class ALMClient:
    """RTC work item access. Stateless apart from the session; results are returned, not stored."""

    def __init__(self, session, cache_file=None):
        if session is None:
            raise RuntimeError("An authenticated RTC session is required")
        self.session = session
        self.cache_file = cache_file or TIMELINE_CACHE_PATH

    # ---- timelines / iterations ----
    def list_timelines(self):
        """[(label, url)] — labels only (url None) when a local cache exists."""
        local_cache = load_local_cache(self.cache_file)
        if local_cache:
            return [(label, None) for label in local_cache]
        url = f"{SERVER}/ccm/process/project-areas/{PROJECT_AREA}/timelines"
        root = xml_root(self.session.get(url, verify=False).text)
        return [(xml_text(tl, "label"), xml_text(tl, "url")) for tl in root.findall(".//{*}timeline")]

    def _fetch_iterations(self, url, depth=0, max_depth=0):
        root = xml_root(self.session.get(url, verify=False).text)
        iterations = []
        for it in root.findall(".//{*}iteration"):
            child = xml_text(it, "childern-iterations-url") or xml_text(it, "children-iterations-url")
            children = []
            if child and depth < max_depth:
                children = self._fetch_iterations(child, depth + 1, max_depth)
            iterations.append({"label": xml_text(it, "label"), "url": xml_text(it, "url"),
                               "children": children, "child_url": child})
        return iterations

    def refresh_timeline_cache(self):
        """Fetch all timelines + first iteration layer and overwrite the local cache."""
        url = f"{SERVER}/ccm/process/project-areas/{PROJECT_AREA}/timelines"
        root = xml_root(self.session.get(url, verify=False).text)
        timeline_cache = {}
        for tl in root.findall(".//{*}timeline"):
            timeline_cache[xml_text(tl, "label")] = self._fetch_iterations(xml_text(tl, "url") + "/iterations")
        save_local_cache(timeline_cache, self.cache_file)
        return timeline_cache

    def refresh_single_timeline(self, timeline_label):
        url = f"{SERVER}/ccm/process/project-areas/{PROJECT_AREA}/timelines"
        root = xml_root(self.session.get(url, verify=False).text)
        for tl in root.findall(".//{*}timeline"):
            if xml_text(tl, "label") == timeline_label:
                data = self._fetch_iterations(xml_text(tl, "url") + "/iterations")
                cache = load_local_cache(self.cache_file)
                cache[timeline_label] = data
                save_local_cache(cache, self.cache_file)
                return data
        return []

    def fetch_children_for_node(self, child_url):
        """One layer of child iterations for lazy loading."""
        if not child_url:
            return []
        return self._fetch_iterations(child_url)

    def update_cache_with_children(self, timeline, chain, child_nodes):
        """Write lazily loaded children back into the timeline cache."""
        try:
            cache = load_local_cache(self.cache_file)
            level = cache.get(timeline)
            if not level:
                return
            target = None
            for label in chain:
                found = next((n for n in level if str(n.get("label") or "").strip() == label), None)
                if not found:
                    return
                target = found
                level = found.get("children") or []
            if target is None:
                return
            target["children"] = child_nodes
            save_local_cache(cache, self.cache_file)
        except Exception as e:
            logger.warning(f"Failed to update ALM cache: {e}")

    def list_scopes(self):
        return flatten_alm_scope_options(load_local_cache(self.cache_file))

    # ---- work items ----
    def parse_defect_data(self, d, iteration_name):
        wi_id = d.get(PROP_ID)
        attr = lambda key: resolve_simple_attr(self.session, (d.get(key) or {}).get("rdf:resource"))
        return {
            "ID": wi_id,
            "Comment": "",
            "Type": attr(PROP_TYPE),
            "Title": d.get(PROP_TITLE, ""),
            "State": attr(PROP_STATE),
            "Resolution": attr(PROP_RESOLUTION),
            "Severity": attr(PROP_SEVERITY),
            "Planned For": iteration_name,
            "Filed Against": attr(PROP_FILED_AGAINST),
            "Approvals": resolve_approvals_via_rpt(self.session, wi_id),
            "Owned By": attr("rtc_cm:ownedBy"),
            "Issue Reason": "",
        }

    def fetch_iteration_data(self, iteration_name, iteration_url,
                             progress_cb: Optional[Callable[[str, int, int, str], None]] = None):
        """Query + parse + quality-check one iteration. Returns (all_defects, qc_issues).
        progress_cb(stage, current, total, message), stage in {"query", "parse", "qc"}."""
        planned_for_id = iteration_url.rsplit("/", 1)[-1]
        t0 = time.time()
        if progress_cb:
            progress_cb("query", 0, 0, f"Querying defect data for {iteration_name}...")

        defects = rtc_query_defects_huge(self.session, planned_for_id)
        t1 = time.time()
        if progress_cb:
            progress_cb("parse", 0, len(defects),
                        f"Query complete, {len(defects)} records, starting parsing (7 network requests per record)...")

        parsed = []
        with ThreadPoolExecutor(max_workers=20) as executor:
            futures = [executor.submit(self.parse_defect_data, d, iteration_name) for d in defects]
            for i, future in enumerate(as_completed(futures), 1):
                try:
                    result = future.result(timeout=60)
                except Exception as e:
                    logger.warning(f"Failed to parse record {i}: {e}")
                    result = None
                if result:
                    parsed.append(result)
                if i % 10 == 0 or i == len(defects):
                    elapsed = time.time() - t1
                    speed = i / elapsed if elapsed > 0 else 0
                    eta = (len(defects) - i) / speed if speed > 0 else 0
                    if progress_cb:
                        progress_cb("parse", i, len(defects),
                                    f"Parse progress {i}/{len(defects)} ({speed:.1f} records/s, ETA {eta:.0f}s)")

        if progress_cb:
            progress_cb("qc", 0, 0, "Running QC analysis...")
        issues = alm_quality_check(parsed)
        logger.info(f"{iteration_name}: {len(parsed)} records, {len(issues)} issues, total {time.time() - t0:.1f}s")
        return parsed, issues


# ============================================================================
# 6. QC (TestInstances / Requirements) analysis
# ============================================================================
def find_qc_csv(base_dir):
    for root, _, files in os.walk(base_dir):
        if "QC_TestInstances.csv" in files and "QC" in root:
            return os.path.join(root, "QC_TestInstances.csv")
    raise FileNotFoundError("QC_TestInstances.csv not found")


def find_qc_requirements_csv(base_dir):
    for root, _, files in os.walk(base_dir):
        if "QC_Requirements.csv" in files and "QC" in root:
            return os.path.join(root, "QC_Requirements.csv")
    raise FileNotFoundError("QC_Requirements.csv not found")


def extract_baseline_from_path(path):
    for part in path.split("\\"):
        m = re.search(r'baseline\s*[_ ]*([0-9.]+)', part, re.IGNORECASE)
        if m:
            return m.group(1)
    return None


def build_qc_scope_tree(paths):
    """Nested folder tree {label: {"path": "\\L1\\L2", "children": {...}}} from folder paths."""
    tree = {}
    for path in paths:
        current_path = ""
        cursor = tree
        for part in [p for p in path.split("\\") if p]:
            current_path = current_path + "\\" + part
            node = cursor.setdefault(part, {"path": current_path, "children": {}})
            cursor = node["children"]
    return tree


def collapse_qc_paths(paths):
    """Drop paths already covered by a selected parent path."""
    collapsed = []
    for path in sorted(set(paths), key=lambda p: (p.count("\\"), p.lower())):
        if any(path == kept or path.startswith(kept + "\\") for kept in collapsed):
            continue
        collapsed.append(path)
    return collapsed


@dataclass
class QCResult:
    prefixes: List[str] = field(default_factory=list)
    baselines: List[str] = field(default_factory=list)
    filtered_rows: List[Dict[str, str]] = field(default_factory=list)   # TestInstances in scope
    problem_rows: List[Dict[str, str]] = field(default_factory=list)    # non-Passed and no RB_Comment
    passed_count: int = 0
    status_counts: Dict[str, int] = field(default_factory=dict)         # "<status> with/without RB_Comment"
    filtered_req_df: Optional[pd.DataFrame] = None                      # Requirement_Type == 105 in baselines
    req_status_counts: Dict[str, int] = field(default_factory=dict)     # Direct_Cover_Status


class QCData:
    """QC_TestInstances.csv / QC_Requirements.csv of a downloaded build."""

    def __init__(self):
        self.rows = []
        self.folder_paths = set()
        self.req_df = None

    def load(self, base_dir):
        """Raises FileNotFoundError when QC_TestInstances.csv is missing."""
        csv_path = find_qc_csv(base_dir)
        with open(csv_path, encoding=QC_ENCODING, newline="") as f:
            self.rows = list(csv.DictReader(f, delimiter=QC_CSV_DELIMITER, quotechar='"'))
        self.folder_paths = {r[FOLDER_COL] for r in self.rows if r.get(FOLDER_COL)}
        try:
            df = pd.read_csv(find_qc_requirements_csv(base_dir), sep=";", encoding=QC_ENCODING,
                             quotechar='"', dtype=str, keep_default_na=False, on_bad_lines="skip").fillna("")
            df.columns = [c.strip().strip('"') for c in df.columns]
            self.req_df = df
        except FileNotFoundError:
            self.req_df = None

    def clear(self):
        self.rows = []
        self.folder_paths = set()
        self.req_df = None

    def available_folder_paths(self):
        return sorted(self.folder_paths, key=lambda path: (path.count("\\"), path.lower()))

    def analyze(self, prefixes):
        """Analyze TestInstances (and Requirements of the same baselines) under the folder prefixes."""
        if not prefixes:
            raise ValueError("Please select a Folder")

        ti_rows = [r for r in self.rows
                   if r.get(FOLDER_COL) and any(r[FOLDER_COL].startswith(p) for p in prefixes)]
        baselines = sorted({b for b in (extract_baseline_from_path(r[FOLDER_COL]) for r in ti_rows) if b})

        passed_count = 0
        stat_counter = Counter()
        problem_rows = []
        for r in ti_rows:
            status = (r.get(RUN_STATUS_COL) or "").strip()
            comment = (r.get("RB_Comment") or "").strip()
            if status == "Passed":
                passed_count += 1
            else:
                stat_counter[f"{status} {'with' if comment else 'without'} RB_Comment"] += 1
                if not comment:
                    problem_rows.append(r)

        result = QCResult(prefixes=list(prefixes), baselines=baselines, filtered_rows=list(ti_rows),
                          problem_rows=problem_rows, passed_count=passed_count,
                          status_counts=dict(sorted(stat_counter.items())))

        if self.req_df is not None and baselines:
            req_df = self.req_df[(self.req_df["RB_Baseline"].isin(baselines)) &
                                 (self.req_df["Requirement_Type"] == "105")].copy()
            result.filtered_req_df = req_df
            result.req_status_counts = dict(Counter(req_df["Direct_Cover_Status"]))
        return result


# ============================================================================
# 7. Excel report export
# ============================================================================
def _excel_modules():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    return Workbook, Alignment, Border, Font, PatternFill, Side, get_column_letter


AS_COLUMNS = ("File_name", "Object_Text", "URL", "RB_Baseline",
              "RB_Top_FNID", "Object_Identifier", "RB_AS_Status",
              "RB_Change_Required", "RB_Agreed_by_Customer",
              "Issue Reason", "Corresponding ALM", "Corresponding SWXS",
              "RB_Planning_Feature", "RB_Allocation")
XS_COLUMNS = ("File_name", "Object_Text", "URL", "RB_Baseline",
              "RB_RS_CP_Status", "RB_TestStatus", "RB_ReviewStatus",
              "Issue Reason", "Story Description", "Story_URL")
ALM_COLUMNS = ["ID", "Comment", "Type", "Title", "State", "Resolution",
               "Severity", "Planned For", "Filed Against",
               "Approvals", "Owned By", "Issue Reason"]
QC_TI_COLUMNS = ["TestSetFolder_Path_full", "TestSet_Name", "TestExec_Status", "Test_Name", "RB_Comment"]
QC_REQ_COLUMNS = ["RB_Baseline", "Direct_Cover_Status", "ReqFolder_Path", "Name", "Author", "Requirement_Type"]


class _Styles:
    def __init__(self):
        _, Alignment, _, Font, PatternFill, _, _ = _excel_modules()
        self.h_fill = PatternFill(start_color="1B2A38", end_color="1B2A38", fill_type="solid")
        self.h_font = Font(name="Segoe UI", color="FFFFFF", bold=True, size=10)
        self.p_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
        self.p_font = Font(name="Segoe UI", color="922B21", size=9)
        self.ok_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
        self.ok_font = Font(name="Segoe UI", color="1E6B35", size=9)
        self.align = Alignment(horizontal="left", vertical="center")
        self.h_align = Alignment(horizontal="center", vertical="center")


def _write_header(ws, cols, row, st):
    _, _, _, _, _, _, get_column_letter = _excel_modules()
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = st.h_fill
        cell.font = st.h_font
        cell.alignment = st.h_align
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{row}"
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _auto_width(ws, mn=10, mx=50):
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, mn), mx)


def _write_summary_block(ws, summary_lines, start_row=1):
    """Collapsible (hidden by default) summary block; returns the next free row."""
    _, _, _, Font, _, _, _ = _excel_modules()
    ws.sheet_properties.outlinePr.summaryBelow = False
    r = start_row
    for line in summary_lines:
        cell = ws.cell(row=r, column=1, value=line)
        if line and not line.startswith(" "):
            cell.font = Font(name="Segoe UI", size=10, bold=True)
        else:
            cell.font = Font(name="Segoe UI", size=9)
        ws.row_dimensions[r].outline_level = 1
        ws.row_dimensions[r].hidden = True
        r += 1
    ws.row_dimensions[r].outline_level = 0
    return r + 1


def _write_table_sheet(wb, title, summary_lines, cols, rows, is_problem, st):
    """rows: list of dict; problem rows first."""
    ws = wb.create_sheet(title)
    next_r = _write_summary_block(ws, summary_lines, start_row=1)
    _write_header(ws, cols, next_r, st)
    ri = next_r + 1
    for row in sorted(rows, key=lambda r: 0 if is_problem(r) else 1):
        prob = is_problem(row)
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=ri, column=ci, value=row.get(c, ""))
            cell.font = st.p_font if prob else st.ok_font
            cell.fill = st.p_fill if prob else st.ok_fill
            cell.alignment = st.align
        ri += 1
    _auto_width(ws)
    return ws


def _has_issue(row):
    return bool(str(row.get("Issue Reason", "")).strip())


def _qc_is_problem(rw):
    return (rw.get(RUN_STATUS_COL) or "").strip().lower() != "passed" and not (rw.get("RB_Comment") or "").strip()


def _split_l1_l2(rw):
    parts = [p for p in (rw.get(FOLDER_COL) or "").strip().split("\\") if p]
    return (parts[0] if parts else "(unclassified)",
            parts[1] if len(parts) >= 2 else "(unclassified)")


def _write_dashboard(ws, cs, qc, qc_selected_paths):
    """Technical dashboard (left: 10_AS + 40_Gen_SWXS, right: ALM + QC) starting at row 6."""
    _, Alignment, Border, Font, PatternFill, Side, _ = _excel_modules()

    CLR_G, CLR_R, CLR_K, CLR_B = "1E6B35", "922B21", "999999", "366092"
    C_SEC = Font(name="Segoe UI", bold=True, size=10, color=CLR_B)
    F_G = Font(name="Segoe UI", size=9, color=CLR_G, bold=True)
    F_R = Font(name="Segoe UI", size=9, color=CLR_R, bold=True)
    F_K = Font(name="Segoe UI", size=9, color=CLR_K, bold=True)
    F_GN = Font(name="Segoe UI", size=9, color=CLR_G)
    F_RN = Font(name="Segoe UI", size=9, color=CLR_R)
    F_KN = Font(name="Segoe UI", size=9, color=CLR_K)
    RALIGN = Alignment(horizontal="right", vertical="center", wrap_text=False)
    LALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
    CALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    FILL_HDR = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    FILL_LBL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    FILL_SUB = PatternFill(start_color="FAFAFA", end_color="FAFAFA", fill_type="solid")
    thin = Side(style="thin", color="BFBFBF")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    def wr(row, col, val, font=None, align=None, fill=None):
        cell = ws.cell(row=row, column=col, value=val)
        if font:
            cell.font = font
        if align:
            cell.alignment = align
        if fill:
            cell.fill = fill
        cell.border = bd

    def pair(row, c_lbl, c_val, label, val, f_lbl, f_val, fill=FILL_LBL):
        wr(row, c_lbl, label, f_lbl, LALIGN, fill)
        wr(row, c_val, val, f_val, RALIGN, fill)

    def section(row, col, title):
        wr(row, col, title, C_SEC, CALIGN, FILL_HDR)
        wr(row, col + 1, "", C_SEC, CALIGN, FILL_HDR)

    # ---- left column (A/B) ----
    cl = 6
    section(cl, 1, "■ 10_AS Status Distribution"); cl += 1
    ss = cs["as_status"]
    gk = list(ss["groups"].keys())
    pair(cl, 1, 2, "Total", ss["total"], F_K, F_K); cl += 1
    pair(cl, 1, 2, gk[3], ss["groups"][gk[3]], F_R, F_R); cl += 1
    pair(cl, 1, 2, gk[1], ss["groups"][gk[1]], F_K, F_K); cl += 1
    pair(cl, 1, 2, gk[0], ss["groups"][gk[0]], F_G, F_G); cl += 1
    pair(cl, 1, 2, "  change=yes", ss["accepted_change_yes"], F_GN, F_G, FILL_SUB); cl += 1
    pair(cl, 1, 2, "    Has ALM assignment", f'{ss["acc_yes_alm_has"]}  ({ss["acc_yes_alm_has_pct"]})', F_GN, F_G, FILL_SUB); cl += 1
    pair(cl, 1, 2, "    No ALM assignment", f'{ss["acc_yes_alm_none"]}  ({ss["acc_yes_alm_none_pct"]})', F_RN, F_R, FILL_SUB); cl += 1
    pair(cl, 1, 2, "    Has SWXS link", f'{ss["acc_yes_swxs_has"]}  ({ss["acc_yes_swxs_has_pct"]})', F_GN, F_G, FILL_SUB); cl += 1
    pair(cl, 1, 2, "    No SWXS link", f'{ss["acc_yes_swxs_none"]}  ({ss["acc_yes_swxs_none_pct"]})', F_RN, F_R, FILL_SUB); cl += 1
    pair(cl, 1, 2, "  change≠yes", ss["accepted_change_no"], F_KN, F_K, FILL_SUB); cl += 1
    pair(cl, 1, 2, gk[2], ss["groups"][gk[2]], F_K, F_K); cl += 1
    cl += 1

    section(cl, 1, "■ 40_Gen_SWXS Status Distribution"); cl += 1
    xs = cs["xs_status"]
    pair(cl, 1, 2, "Total", xs["total"], F_K, F_K); cl += 1
    pair(cl, 1, 2, "accepted/validated", xs["accepted_validated"], F_G, F_G); cl += 1
    pair(cl, 1, 2, "  Has Story link", f'{xs["story_has"]}  ({xs["story_has_pct"]})', F_GN, F_G, FILL_SUB); cl += 1
    pair(cl, 1, 2, "  No Story link", f'{xs["story_none"]}  ({xs["story_none_pct"]})', F_RN, F_R, FILL_SUB); cl += 1
    pair(cl, 1, 2, "  Test passed", f'{xs["test_pass"]}  ({xs["test_pass_pct"]})', F_GN, F_G, FILL_SUB); cl += 1
    pair(cl, 1, 2, "  Test not passed", f'{xs["test_fail"]}  ({xs["test_fail_pct"]})', F_RN, F_R, FILL_SUB); cl += 1
    pair(cl, 1, 2, "  Review ok", f'{xs["review_ok"]}  ({xs["review_ok_pct"]})', F_GN, F_G, FILL_SUB); cl += 1
    pair(cl, 1, 2, "  Review not ok", f'{xs["review_no"]}  ({xs["review_no_pct"]})', F_RN, F_R, FILL_SUB); cl += 1

    # ---- right column (F/G) ----
    cr = 6
    section(cr, 6, "■ ALM Status Distribution"); cr += 1
    al = cs["alm_status"]
    pair(cr, 6, 7, "Total", al["total"], F_K, F_K); cr += 1
    pair(cr, 6, 7, "  Closed", f'{al["closed"]}  ({al["closed_pct"]})', F_GN, F_G, FILL_SUB); cr += 1
    pair(cr, 6, 7, "  Not Closed", f'{al["not_closed"]}  ({al["not_closed_pct"]})', F_RN, F_R, FILL_SUB); cr += 1
    for t in al["type_order"]:
        cnt = al["type_counts"][t]
        pair(cr, 6, 7, t, cnt, F_K, F_K); cr += 1
        awm = (al.get("type_awm") or {}).get(t)
        if awm:
            pair(cr, 6, 7, "  AWM assessment done", f'{awm["done"]}  ({_pct(awm["done"], cnt)})', F_GN, F_G, FILL_SUB); cr += 1
            pair(cr, 6, 7, "  AWM assessment not done", f'{awm["not_done"]}  ({_pct(awm["not_done"], cnt)})', F_RN, F_R, FILL_SUB); cr += 1
    cr += 1

    # QC TestInstance: first-level baseline -> second-level folder -> status -> RB_Comment
    section(cr, 6, "■ QC TestInstance Analysis Summary"); cr += 1
    qc_rows = list(qc.filtered_rows) if qc else []
    qc_total = len(qc_rows)

    sel_pairs = set()
    for sp in (qc_selected_paths or []):
        parts = [p for p in sp.split("\\") if p]
        if len(parts) >= 2:
            sel_pairs.add((parts[0], parts[1]))
        elif len(parts) == 1:
            sel_pairs.add((parts[0], None))

    def row_selected(l1, l2):
        if not sel_pairs:
            return True
        return any(sl1 == l1 and (sl2 is None or sl2 == l2) for sl1, sl2 in sel_pairs)

    l1_order, l1_map = [], {}
    for rw in qc_rows:
        l1, l2 = _split_l1_l2(rw)
        if not row_selected(l1, l2):
            continue
        if l1 not in l1_map:
            l1_map[l1] = {"l2_order": [], "l2_map": {}}
            l1_order.append(l1)
        entry = l1_map[l1]
        if l2 not in entry["l2_map"]:
            entry["l2_map"][l2] = []
            entry["l2_order"].append(l2)
        entry["l2_map"][l2].append(rw)

    pair(cr, 6, 7, "Total", qc_total, F_K, F_K); cr += 1
    if qc_total == 0 or not l1_order:
        pair(cr, 6, 7, "(no QC data)", "", F_KN, F_KN); cr += 1
    for l1 in l1_order:
        entry = l1_map[l1]
        l1_cnt = sum(len(entry["l2_map"][l2]) for l2 in entry["l2_order"])
        pair(cr, 6, 7, f"■ {l1}", f"{l1_cnt}  ({_pct(l1_cnt, qc_total)})", F_K, F_K, FILL_HDR); cr += 1
        for l2 in entry["l2_order"]:
            l2_rows = entry["l2_map"][l2]
            l2_cnt = len(l2_rows)
            pair(cr, 6, 7, f"  ▸ {l2}", f"{l2_cnt}  ({_pct(l2_cnt, l1_cnt)})", F_K, F_K, FILL_LBL); cr += 1
            st_order, st_map = [], {}
            for rw in l2_rows:
                s = (rw.get(RUN_STATUS_COL) or "").strip() or "(empty)"
                if s not in st_map:
                    st_map[s] = []
                    st_order.append(s)
                st_map[s].append(rw)
            for s in st_order:
                rows = st_map[s]
                cnt = len(rows)
                passed = s.lower() == "passed"
                f = F_G if passed else F_K
                pair(cr, 6, 7, f"    {s}", f"{cnt}  ({_pct(cnt, l2_cnt)})", f, f, FILL_SUB); cr += 1
                if not passed:
                    has_c = sum(1 for rw in rows if (rw.get("RB_Comment") or "").strip())
                    pair(cr, 6, 7, "      Has RB_Comment", f"{has_c}  ({_pct(has_c, cnt)})", F_GN, F_G, FILL_SUB); cr += 1
                    pair(cr, 6, 7, "      No RB_Comment", f"{cnt - has_c}  ({_pct(cnt - has_c, cnt)})", F_RN, F_R, FILL_SUB); cr += 1

    # QC Requirement: RB_Baseline -> Direct_Cover_Status
    section(cr, 6, "■ QC Requirement Analysis Summary"); cr += 1
    req_df = qc.filtered_req_df if qc else None
    if req_df is not None and not req_df.empty:
        req_total = len(req_df)
        pair(cr, 6, 7, "Total", req_total, F_K, F_K); cr += 1
        baselines = list(req_df["RB_Baseline"].dropna().unique()) if "RB_Baseline" in req_df.columns else []
        for bl in baselines:
            bl_rows = req_df[req_df["RB_Baseline"] == bl]
            bl_cnt = len(bl_rows)
            pair(cr, 6, 7, f"■ Baseline {bl}", f"{bl_cnt}  ({_pct(bl_cnt, req_total)})", F_K, F_K, FILL_HDR); cr += 1
            counts = bl_rows["Direct_Cover_Status"].value_counts() if "Direct_Cover_Status" in bl_rows.columns else {}
            for s, cnt in counts.items():
                f = F_G if str(s).strip().lower() == "passed" else F_K
                pair(cr, 6, 7, f"  {s}", f"{cnt}  ({_pct(cnt, bl_cnt)})", f, f, FILL_SUB); cr += 1
    else:
        pair(cr, 6, 7, "(no Requirement data)", "", F_KN, F_KN); cr += 1

    for letter, width in (("A", 42), ("B", 16), ("C", 2), ("D", 2), ("E", 2), ("F", 42), ("G", 16)):
        ws.column_dimensions[letter].width = width


def _write_selection_summary(ws, selection_lines):
    _, _, _, Font, _, _, _ = _excel_modules()
    row = ws.max_row + 2
    ws.cell(row=row, column=1, value="User Input Selection Summary").font = Font(name="Segoe UI", bold=True, size=14, color="1B2A38")
    row += 1
    for line in selection_lines:
        cell = ws.cell(row=row, column=1, value=line)
        if line.startswith("["):
            cell.font = Font(name="Segoe UI", size=10, bold=True, color="366092")
        elif line.startswith("    ✔"):
            cell.font = Font(name="Segoe UI", size=9, color="1E6B35")
        elif line.startswith("    ✗"):
            cell.font = Font(name="Segoe UI", size=9, color="999999")
        elif line.startswith("  "):
            cell.font = Font(name="Segoe UI", size=9, bold=True)
        else:
            cell.font = Font(name="Segoe UI", size=9)
        row += 1


def export_all_reports(file_path, result, project="", build="", build_date="",
                       selection_lines=None, qc_selected_paths=None, log: LogFunc = _noop_log):
    """Write the combined workbook: Summary, 10_AS, 40_Gen_SWXS, ALM, QC TestInstance, QC Requirement.
    `result` is an AnalysisResult."""
    Workbook, _, _, Font, _, _, _ = _excel_modules()
    st = _Styles()
    wb = Workbook()

    # ---- Summary ----
    ws = wb.active
    ws.title = "Summary"
    log("Writing summary sheet...", "info")
    ws["A1"] = "Technical Dashboard"
    ws["A1"].font = Font(name="Segoe UI", bold=True, size=16)
    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A3"] = f"Project: {project}"
    ws["A4"] = f"Build: {build}  ({build_date})"
    try:
        cs = build_cross_ref_stats(result.as_rows, result.xs_rows, result.alm_defects)
        _write_dashboard(ws, cs, result.qc, qc_selected_paths)
    except Exception as e:
        logger.warning(f"Cross-reference stats write failed: {e}")
    if selection_lines:
        _write_selection_summary(ws, selection_lines)
    _auto_width(ws)

    # ---- 10_AS / 40_Gen_SWXS ----
    log("Writing 10_AS Report sheet...", "info")
    _write_table_sheet(wb, "10_AS Report", build_as_summary(result.as_rows), AS_COLUMNS,
                       result.as_rows, _has_issue, st)
    log("Writing 40_Gen_SWXS Report sheet...", "info")
    _write_table_sheet(wb, "40_Gen_SWXS Report", build_xs_summary(result.xs_rows), XS_COLUMNS,
                       result.xs_rows, _has_issue, st)

    # ---- ALM ----
    if result.alm_defects:
        log("Writing ALM sheet...", "info")
        issue_map = {row.get("ID", ""): row.get("Issue Reason", "") for row in result.alm_issues}
        alm_rows = []
        for d in result.alm_defects:
            row = dict(d)
            row["Issue Reason"] = issue_map.get(d.get("ID", ""), "")
            alm_rows.append(row)
        alm_prob = sum(1 for r in alm_rows if r["Issue Reason"].strip())
        reason_counter = Counter(r["Issue Reason"].strip() for r in alm_rows if r["Issue Reason"].strip())
        lines = ["ALM Analysis Summary",
                 f"Total records: {len(alm_rows)}  |  Issues: {alm_prob}"]
        if reason_counter:
            lines.append("Issues: " + "  |  ".join(
                f"{r}: {c}" for r, c in sorted(reason_counter.items(), key=lambda x: (-x[1], x[0]))))
        else:
            lines.append("Issues: None")
        lines.append("")
        _write_table_sheet(wb, "ALM Analysis", lines, ALM_COLUMNS, alm_rows, _has_issue, st)

    # ---- QC TestInstance ----
    qc = result.qc
    try:
        qc_rows = list(qc.filtered_rows) if qc else []
        lines = ["QC TestInstance Analysis Summary",
                 f"Total records: {len(qc_rows)}  |  Issues: {len(qc.problem_rows) if qc else 0}"]
        counter = {}
        for rw in qc_rows:
            l1, l2 = _split_l1_l2(rw)
            info = counter.setdefault(l1, {}).setdefault(l2, {"total": 0, "status": {}})
            info["total"] += 1
            s = (rw.get(RUN_STATUS_COL) or "").strip() or "(empty)"
            info["status"][s] = info["status"].get(s, 0) + 1
        if counter:
            for l1 in sorted(counter):
                lines.append(f"  {l1} (total {sum(v['total'] for v in counter[l1].values())})")
                for l2 in sorted(counter[l1]):
                    info = counter[l1][l2]
                    sts = [f"{s}: {c}" for s, c in sorted(info["status"].items(), key=lambda x: (-x[1], x[0]))]
                    lines.append(f"    {l2} (total {info['total']}): " + "  |  ".join(sts))
        else:
            lines.append("Status distribution: None")
        lines.append("")
        log("Writing QC sheet...", "info")
        _write_table_sheet(wb, "QC TestInstance", lines, QC_TI_COLUMNS, qc_rows, _qc_is_problem, st)
    except Exception as e:
        log(f"QC sheet write failed: {e}", "warning")

    # ---- QC Requirement ----
    try:
        req_df = qc.filtered_req_df if qc else None
        if req_df is not None and not req_df.empty:
            log("Writing QC Requirement sheet...", "info")
            req_rows = req_df.to_dict("records")
            is_prob = lambda r: str(r.get("Direct_Cover_Status", "")).strip() != "Passed"
            counter = {}
            for r in req_rows:
                bl = str(r.get("RB_Baseline", "")).strip() or "(unclassified)"
                s = str(r.get("Direct_Cover_Status", "")).strip() or "(empty)"
                counter.setdefault(bl, {})
                counter[bl][s] = counter[bl].get(s, 0) + 1
            lines = ["QC Requirement Analysis Summary",
                     f"Total records: {len(req_rows)}  |  Issues: {sum(1 for r in req_rows if is_prob(r))}"]
            for bl in sorted(counter):
                lines.append(f"  Baseline {bl} (total {sum(counter[bl].values())})")
                lines.append("    " + "  |  ".join(
                    f"{s}: {c}" for s, c in sorted(counter[bl].items(), key=lambda x: (-x[1], x[0]))))
            lines.append("")
            _write_table_sheet(wb, "QC Requirement", lines, QC_REQ_COLUMNS, req_rows, is_prob, st)
    except Exception as e:
        log(f"QC Requirement sheet write failed: {e}", "warning")

    wb.save(file_path)
    log(f"Report exported successfully: {file_path}", "success")
    return file_path


def export_as_rows_to_excel(rows, file_path):
    """Stand-alone 10_AS export (rows as returned by DoorsDataSession.analyze_as)."""
    Workbook, Alignment, _, Font, PatternFill, _, get_column_letter = _excel_modules()
    wb = Workbook()
    ws = wb.active
    ws.title = "10_AS Analysis"
    h_fill = PatternFill(start_color="1B2A38", end_color="1B2A38", fill_type="solid")
    h_font = Font(name="Segoe UI", color="FFFFFF", bold=True)
    p_font = Font(name="Segoe UI", color="C0392B", size=9)
    n_font = Font(name="Segoe UI", size=9)
    align = Alignment(horizontal="left", vertical="center")
    for i, c in enumerate(AS_COLUMNS, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.fill, cell.font = h_fill, h_font
    ws.auto_filter.ref = f"A1:{get_column_letter(len(AS_COLUMNS))}1"
    ws.freeze_panes = ws.cell(row=2, column=1)
    for ri, row in enumerate(rows, 2):
        font = p_font if _has_issue(row) else n_font
        for ci, c in enumerate(AS_COLUMNS, 1):
            cell = ws.cell(row=ri, column=ci, value=row.get(c, ""))
            cell.font, cell.alignment = font, align
    for i in range(1, len(AS_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    wb.save(file_path)
    return file_path


# ============================================================================
# 8. AnalysisBackend facade (full analysis pipeline)
# ============================================================================
@dataclass
class AnalysisRequest:
    project: str = ""
    build: str = ""
    as_file_names: List[str] = field(default_factory=list)
    as_features: List[str] = field(default_factory=list)
    xs_file_names: List[str] = field(default_factory=list)
    xs_features: List[str] = field(default_factory=list)
    alm_scopes: List[Any] = field(default_factory=list)   # scope dicts or "display" strings
    qc_paths: List[str] = field(default_factory=list)


@dataclass
class AnalysisResult:
    as_rows: List[Dict[str, Any]] = field(default_factory=list)
    xs_rows: List[Dict[str, Any]] = field(default_factory=list)
    alm_defects: List[Dict[str, Any]] = field(default_factory=list)
    alm_issues: List[Dict[str, Any]] = field(default_factory=list)
    qc: Optional[QCResult] = None
    cross_ref_stats: Dict[str, Any] = field(default_factory=dict)

    @property
    def as_summary(self):
        return build_as_summary(self.as_rows)

    @property
    def xs_summary(self):
        return build_xs_summary(self.xs_rows)

    @property
    def alm_summary(self):
        return build_alm_summary_text(self.alm_defects, self.alm_issues)


class AnalysisBackend:
    """Owns the Jenkins/DOORS data, optional ALM client and QC data, and runs the whole pipeline."""

    TOTAL_STEPS = 5

    def __init__(self, log: LogFunc = _noop_log):
        self.log = log
        self.data = DoorsDataSession()
        self.qc = QCData()
        self.alm: Optional[ALMClient] = None
        self.force_alm_refresh = False  # ignore ALM cache on next run (build changed / scope changed)

    # ---- connections ----
    def connect_alm(self, username=None, password=None, cred_file=None):
        session = rtc_setup_session(username, password, cred_file or CRED_PATH)
        self.alm = ALMClient(session)
        return self.alm

    @staticmethod
    def jenkins_projects():
        return get_all_projects()

    @staticmethod
    def jenkins_status(project):
        return get_jenkins_status(project)

    # ---- data loading ----
    def switch_project(self):
        """Drop all loaded data (call when the project changes)."""
        self.data.clear()
        self.qc.clear()

    def load_build(self, project, build):
        """Download + scan a build and load its QC CSVs. Returns (as_files, xs_files)."""
        as_files, xs_files = self.data.download_and_scan(project, build, self.log)
        self.force_alm_refresh = True  # new build: re-fetch ALM data
        self.load_qc()
        return as_files, xs_files

    def load_qc(self):
        try:
            self.qc.load(self.data.base_dir)
            self.log("QC data loaded", "success")
        except Exception as e:
            self.qc.clear()
            self.log(f"QC data load failed: {e}", "warning")

    def alm_scopes(self):
        return self.alm.list_scopes() if self.alm else flatten_alm_scope_options()

    def qc_folder_tree(self):
        return build_qc_scope_tree(self.qc.available_folder_paths())

    # ---- ALM data for several iterations ----
    def fetch_alm(self, scopes, progress_cb: Optional[Callable[[str, int], None]] = None):
        """Fetch (or load from cache) and merge ALM data of the scopes. Returns (defects, issues).
        progress_cb(message, percent_within_step 0..100)."""
        scope_map = {s["display"]: s for s in self.alm_scopes()}
        resolved = []
        for s in scopes:
            scope = scope_map.get(s) if isinstance(s, str) else s
            if not scope:
                raise KeyError(f"Unknown ALM scope: {s}")
            resolved.append(scope)

        merged_defects, merged_issues = [], []
        seen_defects, seen_issues = set(), set()
        alm_cache = load_alm_cache()

        for idx, scope in enumerate(resolved):
            display = scope["display"]
            self.log(f"ALM Scope ({idx + 1}/{len(resolved)}): {display}", "info")
            cached = None if self.force_alm_refresh else alm_cache.get(display)
            if cached:
                self.log(f"{display} (using cache)", "info")
                defects, issues = cached.get("all_defects", []), cached.get("qc_issues", [])
            else:
                if not self.alm:
                    raise RuntimeError("ALM is not connected: call connect_alm() first")

                def on_progress(stage, current, total, message, display=display):
                    pct = 0
                    if stage == "parse" and total > 0:
                        pct = int(current / total * 80)
                    elif stage == "qc":
                        pct = 90
                    self.log(f"    {message}", "info")
                    if progress_cb:
                        progress_cb(f"[{display}] {message}", pct)

                defects, issues = self.alm.fetch_iteration_data(scope["label"], scope["url"], on_progress)
                alm_cache[display] = {"all_defects": defects, "qc_issues": issues}

            for d in defects:
                wid = d.get("ID")
                if wid and wid not in seen_defects:
                    seen_defects.add(wid)
                    merged_defects.append(d)
            for q in issues:
                wid = q.get("ID")
                if wid and wid not in seen_issues:
                    seen_issues.add(wid)
                    merged_issues.append(q)

        save_alm_cache(alm_cache)
        self.log(f"ALM results cached ({len(resolved)} scopes)", "info")
        self.force_alm_refresh = False
        return merged_defects, merged_issues

    # ---- pipeline ----
    def run_full_analysis(self, request: AnalysisRequest, progress: Optional[ProgressFunc] = None):
        """Step 1 load/scan, 2 10_AS, 3 40_Gen_SWXS, 4 ALM, 5 QC. Returns AnalysisResult."""
        total = self.TOTAL_STEPS

        def report(step, message, percent):
            self.log(f"Step {step}/{total}  {message}", "info")
            if progress:
                progress(step, total, message, percent)

        result = AnalysisResult()
        project = str(request.project or "").strip()
        build = str(request.build or "").strip()
        if not build and str(self.data.project or "") == project:
            build = str(self.data.build or "")

        # Step 1
        need_reload = (not self.data.loaded
                       or str(self.data.project or "") != project
                       or str(self.data.build or "") != build)
        if need_reload:
            if not build:
                raise RuntimeError("Jenkins build info is empty for current project, please refresh Jenkins status and retry")
            report(1, "Downloading and scanning Jenkins artifact...", 5)
            self.load_build(project, build)
            report(1, f"Scan complete: {len(self.data.file_map)} AS files, {len(self.data.xs_file_map)} XS files", 15)
        else:
            report(1, "Using currently scanned data", 15)

        # Step 2
        report(2, "Analyzing 10_AS data...", 22)
        if request.as_file_names and request.as_features:
            result.as_rows = self.data.analyze_as(request.as_file_names, request.as_features)
        else:
            self.log("No 10_AS files or baselines selected, skipping 10_AS analysis", "warning")
        report(2, "10_AS analysis complete", 35)

        # Step 3
        report(3, "Analyzing 40_Gen_SWXS data...", 42)
        if request.xs_file_names and request.xs_features:
            result.xs_rows = self.data.analyze_xs(request.xs_file_names, request.xs_features)
        else:
            self.log("No 40_Gen files or baselines selected, skipping XS analysis", "warning")
        report(3, "40_Gen_SWXS analysis complete", 55)

        # Step 4
        report(4, "Fetching ALM data...", 62)
        if request.alm_scopes:
            result.alm_defects, result.alm_issues = self.fetch_alm(
                request.alm_scopes,
                lambda msg, pct: progress and progress(4, total, msg, 62 + int(pct * 0.15)))
        else:
            self.log("No ALM Iteration selected, skipping ALM analysis", "warning")
        report(4, "ALM analysis complete", 80)

        # Step 5
        report(5, "Analyzing QC data...", 88)
        if request.qc_paths:
            if not self.qc.rows:
                self.load_qc()
            if self.qc.rows:
                result.qc = self.qc.analyze(request.qc_paths)
        else:
            self.log("No QC Folder selected, skipping QC analysis", "warning")
        report(5, "QC analysis complete", 100)

        result.cross_ref_stats = build_cross_ref_stats(result.as_rows, result.xs_rows, result.alm_defects)
        return result

    # ---- reporting ----
    def build_selection_summary_lines(self, request: AnalysisRequest):
        """Text block describing available vs. selected inputs (used in the report summary sheet)."""
        lines = []

        def block(title, file_names, features, files_available, feature_pool):
            lines.append(f"[{title} Input Selection]")
            sel_files, sel_feats = set(file_names), set(features)
            lines.append(f"  Files (available {len(files_available)} / selected {len(sel_files)}):")
            lines.extend([f"    {'✔' if f in sel_files else '✗'} {f}" for f in files_available] or ["    (no files available)"])
            lines.append(f"  Baselines (available {len(feature_pool)} / selected {len(sel_feats)}):")
            lines.extend([f"    {'✔' if f in sel_feats else '✗'} {f}" for f in feature_pool] or ["    (no baselines available)"])
            lines.append("")

        as_pool = self.data.list_baselines(request.as_file_names, "as") if request.as_file_names else []
        xs_pool = self.data.list_baselines(request.xs_file_names, "xs") if request.xs_file_names else []
        block("10_AS", request.as_file_names, request.as_features, sorted(self.data.file_map), as_pool)
        block("40_Gen_SWXS", request.xs_file_names, request.xs_features, sorted(self.data.xs_file_map), xs_pool)

        lines.append("[ALM Input Selection]")
        lines.append(f"  Selected Iterations: {len(request.alm_scopes)}")
        for s in request.alm_scopes:
            lines.append(f"    ✔ {s.get('display', '') if isinstance(s, dict) else s}")
        if not request.alm_scopes:
            lines.append("    (no Iteration selected)")
        lines.append("")

        lines.append("[QC Input Selection]")
        lines.append(f"  Selected Folders: {len(request.qc_paths)}")
        for p in request.qc_paths:
            lines.append(f"    ✔ {p}")
        if not request.qc_paths:
            lines.append("    (no Folder selected)")
        lines.append("")
        return lines

    def export_report(self, file_path, result: AnalysisResult, request: AnalysisRequest, build_date=""):
        return export_all_reports(
            file_path, result,
            project=request.project, build=request.build, build_date=build_date,
            selection_lines=self.build_selection_summary_lines(request),
            qc_selected_paths=request.qc_paths,
            log=self.log,
        )
